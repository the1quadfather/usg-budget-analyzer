"""Tie ingested R-1 and DD 1416 amounts to their published references."""

from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
import re
from typing import Literal

import openpyxl
from sqlalchemy import and_, func, select
from sqlalchemy.orm import Session

from parsing.xlsx_ingest import R1XlsxParser
from storage.db import FundingLine, PEExecution, ProgramElement


Stream = Literal["discretionary", "mandatory"]
Basis = Literal["PY Actual", "CY Request", "BY Request"]
TieStatus = Literal["ties", "outside_tolerance", "no_reference"]


@dataclass(frozen=True)
class TieOutRow:
    pb_cycle: int
    fiscal_year: int
    scope: str
    agency: str | None
    stream: Stream
    basis: Basis
    ingested_k: float
    reference_k: float | None
    residual_k: float | None
    residual_pct: float | None
    status: TieStatus
    explanation: str


@dataclass(frozen=True)
class _WorkbookReference:
    fiscal_year: int
    stream: Stream
    basis: Basis
    header: str
    printed_total_k: float | None
    account_totals_k: dict[str, float]


ACCOUNT_TO_AGENCY = {
    "2040A": "Army",
    "1319N": "Navy",
    "3600F": "Air Force",
    "3620F": "Space Force",
    "0400D": "Defense-Wide",
    "0460D": "OT&E",
}
AGENCY_TO_ACCOUNT = {
    agency: account for account, agency in ACCOUNT_TO_AGENCY.items()
}
FUNDING_TYPE_TO_KEY: dict[str, tuple[Stream, Basis]] = {
    "PY Actual": ("discretionary", "PY Actual"),
    "CY Request": ("discretionary", "CY Request"),
    "BY Request": ("discretionary", "BY Request"),
    "PY Mandatory": ("mandatory", "PY Actual"),
    "CY Mandatory": ("mandatory", "CY Request"),
    "BY Mandatory": ("mandatory", "BY Request"),
}
_BASIS_BY_OFFSET: dict[int, Basis] = {
    2: "PY Actual",
    1: "CY Request",
    0: "BY Request",
}
_ACCOUNT_CODE_RE = re.compile(r"^\d{4}[A-Z]$")
_FY_HEADER_RE = re.compile(r"^FY\s*(\d{4})\s*(.*)$", re.IGNORECASE)
_MANDATORY_HEADER_RE = re.compile(
    r"\breconciliation\b|\bmandatory\b|\bpl\s*119", re.IGNORECASE
)
_SUPPLEMENTAL_HEADER_RE = re.compile(r"\bsupplementals?\b", re.IGNORECASE)
_DOWNLOAD_COMMAND = (
    "python acquisition/comptroller_scraper.py --xlsx "
    "--years {pb_cycle} --exhibits rdtee"
)
_XLSX_FIRST_PB_CYCLE = 2012


def _as_amount(value) -> float | None:
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "")
    if not text or text == "-":
        return None
    if text.startswith("(") and text.endswith(")"):
        text = f"-{text[1:-1]}"
    try:
        return float(text)
    except ValueError:
        return None


def _compare(
    ingested_k: float,
    reference_k: float | None,
    tolerance_pct: float,
) -> tuple[float | None, float | None, TieStatus]:
    if tolerance_pct < 0:
        raise ValueError("tolerance_pct must be non-negative")
    if reference_k is None:
        return None, None, "no_reference"

    residual_k = ingested_k - reference_k
    if reference_k == 0:
        residual_pct = 0.0 if residual_k == 0 else None
        status: TieStatus = "ties" if residual_k == 0 else "outside_tolerance"
    else:
        residual_pct = residual_k / reference_k * 100.0
        status = (
            "ties"
            if abs(residual_pct) <= tolerance_pct
            else "outside_tolerance"
        )
    return residual_k, residual_pct, status


def _funding_aggregates(
    session: Session,
) -> dict[tuple[int, int, Stream, Basis, str], float]:
    statement = (
        select(
            FundingLine.pb_cycle,
            FundingLine.fiscal_year,
            FundingLine.funding_type,
            ProgramElement.agency,
            func.sum(FundingLine.amount_thousands),
        )
        .join(ProgramElement, FundingLine.program_element_id == ProgramElement.id)
        .where(FundingLine.pb_cycle.is_not(None))
        .group_by(
            FundingLine.pb_cycle,
            FundingLine.fiscal_year,
            FundingLine.funding_type,
            ProgramElement.agency,
        )
    )
    aggregates: dict[tuple[int, int, Stream, Basis, str], float] = {}
    for pb_cycle, fiscal_year, funding_type, agency, amount in session.execute(
        statement
    ):
        stream_basis = FUNDING_TYPE_TO_KEY.get(funding_type)
        if stream_basis is None:
            continue
        stream, basis = stream_basis
        aggregates[(
            int(pb_cycle), int(fiscal_year), stream, basis, agency,
        )] = float(amount)
    return aggregates


def _amount_columns(
    headers: list[str], pb_cycle: int
) -> dict[int, dict[Stream, int]]:
    columns: dict[int, dict[Stream, int]] = {}
    primary_candidates: defaultdict[int, list[tuple[int, int]]] = defaultdict(
        list
    )
    for column_index, header in enumerate(headers):
        match = _FY_HEADER_RE.match(header)
        if match is None:
            continue
        fiscal_year = int(match.group(1))
        if pb_cycle - fiscal_year not in _BASIS_BY_OFFSET:
            continue
        qualifier = " ".join(match.group(2).lower().rstrip("*").split())
        if _MANDATORY_HEADER_RE.search(qualifier):
            columns.setdefault(fiscal_year, {}).setdefault(
                "mandatory", column_index
            )
            continue
        if (
            _SUPPLEMENTAL_HEADER_RE.search(qualifier)
            and "less supplemental" not in qualifier
        ):
            continue

        if re.search(r"\btotal\b", qualifier):
            priority = 6
        elif re.search(r"\bactuals?\b", qualifier):
            priority = 0
        elif re.search(r"\b(?:discretionary|disc)\b", qualifier) and re.search(
            r"\benacted\b", qualifier
        ):
            priority = 1
        elif re.search(r"\benacted\b", qualifier):
            priority = 2
        elif re.search(r"\b(?:discretionary|disc)\b", qualifier) and re.search(
            r"\brequest\b", qualifier
        ):
            priority = 3
        elif re.search(r"\bbase\b", qualifier):
            priority = 4
        elif re.search(r"\brequest\b", qualifier):
            priority = 5
        else:
            continue
        primary_candidates[fiscal_year].append((priority, column_index))

    for fiscal_year, candidates in primary_candidates.items():
        _priority, column_index = min(candidates)
        columns.setdefault(fiscal_year, {})["discretionary"] = column_index
    return columns


def _read_workbook(
    workbook_path: Path, pb_cycle: int
) -> list[_WorkbookReference]:
    workbook = openpyxl.load_workbook(
        workbook_path, read_only=True, data_only=True
    )
    try:
        worksheet, headers, header_index = R1XlsxParser._find_exhibit_sheet(
            workbook
        )
        amount_columns = _amount_columns(headers, pb_cycle)
        account_index = headers.index("Account")

        total_row = None
        for row in worksheet.iter_rows(
            min_row=1, max_row=header_index, values_only=True
        ):
            if any(
                str(value).strip().lower() == "total of displayed rows"
                for value in row if value is not None
            ):
                total_row = row
                break

        references: list[_WorkbookReference] = []
        for fiscal_year, stream_columns in sorted(amount_columns.items()):
            basis = _BASIS_BY_OFFSET.get(pb_cycle - fiscal_year)
            if basis is None:
                continue
            for stream, column_index in sorted(
                stream_columns.items(), key=lambda item: item[1]
            ):
                account_totals: defaultdict[str, float] = defaultdict(float)
                for row in worksheet.iter_rows(
                    min_row=header_index + 2, values_only=True
                ):
                    if account_index >= len(row) or column_index >= len(row):
                        continue
                    account = str(row[account_index] or "").strip().upper()
                    if not _ACCOUNT_CODE_RE.fullmatch(account):
                        continue
                    amount = _as_amount(row[column_index])
                    if amount is not None:
                        account_totals[account] += amount

                printed_total = (
                    _as_amount(total_row[column_index])
                    if total_row is not None and column_index < len(total_row)
                    else None
                )
                references.append(_WorkbookReference(
                    fiscal_year=fiscal_year,
                    stream=stream,
                    basis=basis,
                    header=headers[column_index],
                    printed_total_k=printed_total,
                    account_totals_k=dict(account_totals),
                ))
        return references
    finally:
        workbook.close()


def _missing_workbook_rows(
    pb_cycle: int,
    aggregates: dict[tuple[int, int, Stream, Basis, str], float],
) -> list[TieOutRow]:
    column_keys = sorted(
        {
            (fiscal_year, stream, basis)
            for cycle, fiscal_year, stream, basis, _agency in aggregates
            if cycle == pb_cycle
        },
        key=lambda key: (key[0], key[2], key[1]),
    )
    explanation = (
        "R-1 workbook is missing; run: "
        + _DOWNLOAD_COMMAND.format(pb_cycle=pb_cycle)
    )
    rows = []
    for fiscal_year, stream, basis in column_keys:
        ingested_k = sum(
            aggregates.get((
                pb_cycle, fiscal_year, stream, basis, agency,
            ), 0.0)
            for agency in AGENCY_TO_ACCOUNT
        )
        rows.append(TieOutRow(
            pb_cycle=pb_cycle,
            fiscal_year=fiscal_year,
            scope="DoD",
            agency=None,
            stream=stream,
            basis=basis,
            ingested_k=ingested_k,
            reference_k=None,
            residual_k=None,
            residual_pct=None,
            status="no_reference",
            explanation=explanation,
        ))
    return rows


def _workbook_rows(
    pb_cycle: int,
    reference: _WorkbookReference,
    aggregates: dict[tuple[int, int, Stream, Basis, str], float],
    tolerance_pct: float,
) -> list[TieOutRow]:
    key = (
        pb_cycle, reference.fiscal_year, reference.stream, reference.basis,
    )
    excluded_reference_k = sum(
        amount
        for account, amount in reference.account_totals_k.items()
        if account not in ACCOUNT_TO_AGENCY
    )
    adjusted_reference_k = (
        reference.printed_total_k - excluded_reference_k
        if reference.printed_total_k is not None
        else None
    )
    ingested_dod_k = sum(
        aggregates.get((*key, agency), 0.0)
        for agency in AGENCY_TO_ACCOUNT
    )
    residual_k, residual_pct, status = _compare(
        ingested_dod_k, adjusted_reference_k, tolerance_pct
    )
    if status == "no_reference":
        explanation = (
            f"workbook column '{reference.header}' has no printed "
            "Total of Displayed Rows"
        )
    elif status == "outside_tolerance":
        explanation = (
            "in-scope ingested funding differs from the workbook after "
            "excluding non-RDT&E accounts; see per-account rows"
        )
    else:
        explanation = ""

    rows = [TieOutRow(
        pb_cycle=pb_cycle,
        fiscal_year=reference.fiscal_year,
        scope="DoD",
        agency=None,
        stream=reference.stream,
        basis=reference.basis,
        ingested_k=ingested_dod_k,
        reference_k=adjusted_reference_k,
        residual_k=residual_k,
        residual_pct=residual_pct,
        status=status,
        explanation=explanation,
    )]

    for account, account_reference_k in sorted(
        reference.account_totals_k.items()
    ):
        agency = ACCOUNT_TO_AGENCY.get(account)
        ingested_k = (
            aggregates.get((*key, agency), 0.0)
            if agency is not None
            else 0.0
        )
        residual_k, residual_pct, status = _compare(
            ingested_k, account_reference_k, tolerance_pct
        )
        if agency is None:
            explanation = "account not ingested (non-RDT&E appropriation)"
        elif status == "outside_tolerance":
            explanation = (
                "ingested funding lines differ from the workbook account total"
            )
        else:
            explanation = ""
        rows.append(TieOutRow(
            pb_cycle=pb_cycle,
            fiscal_year=reference.fiscal_year,
            scope=account,
            agency=agency,
            stream=reference.stream,
            basis=reference.basis,
            ingested_k=ingested_k,
            reference_k=account_reference_k,
            residual_k=residual_k,
            residual_pct=residual_pct,
            status=status,
            explanation=explanation,
        ))
    return rows


def tie_out_r1(
    session: Session,
    raw_dir: Path,
    *,
    tolerance_pct: float = 0.5,
) -> list[TieOutRow]:
    """Compare ingested R-1 totals with workbook grand and account totals."""
    if tolerance_pct < 0:
        raise ValueError("tolerance_pct must be non-negative")
    raw_dir = Path(raw_dir)
    aggregates = _funding_aggregates(session)
    database_cycles = {
        pb_cycle
        for pb_cycle, _fy, _stream, _basis, _agency in aggregates
        if pb_cycle >= _XLSX_FIRST_PB_CYCLE
    }
    workbook_cycles = {
        int(path.parent.parent.name)
        for path in raw_dir.glob("[0-9][0-9][0-9][0-9]/rdtee/fy*_r1.xlsx")
        if path.parent.parent.name.isdigit()
    }

    rows: list[TieOutRow] = []
    for pb_cycle in sorted(database_cycles | workbook_cycles):
        workbook_path = (
            raw_dir / str(pb_cycle) / "rdtee" / f"fy{pb_cycle}_r1.xlsx"
        )
        if not workbook_path.exists():
            rows.extend(_missing_workbook_rows(pb_cycle, aggregates))
            continue
        for reference in _read_workbook(workbook_path, pb_cycle):
            rows.extend(_workbook_rows(
                pb_cycle, reference, aggregates, tolerance_pct
            ))
    return rows


def tie_out_dd1416(
    session: Session,
    *,
    tolerance_pct: float = 0.5,
) -> list[TieOutRow]:
    """Compare latest DD 1416 enacted totals with next-cycle R-1 CY totals."""
    if tolerance_pct < 0:
        raise ValueError("tolerance_pct must be non-negative")
    latest_reports = (
        select(
            PEExecution.fy_start.label("fy_start"),
            PEExecution.agency.label("agency"),
            func.max(PEExecution.report_date).label("report_date"),
        )
        .group_by(PEExecution.fy_start, PEExecution.agency)
        .subquery()
    )
    statement = (
        select(
            PEExecution.fy_start,
            PEExecution.agency,
            latest_reports.c.report_date,
            func.sum(PEExecution.enacted_k),
        )
        .join(
            latest_reports,
            and_(
                PEExecution.fy_start == latest_reports.c.fy_start,
                PEExecution.agency == latest_reports.c.agency,
                PEExecution.report_date == latest_reports.c.report_date,
            ),
        )
        .group_by(
            PEExecution.fy_start,
            PEExecution.agency,
            latest_reports.c.report_date,
        )
        .order_by(PEExecution.fy_start, PEExecution.agency)
    )
    r1_aggregates = _funding_aggregates(session)
    rows = []
    for fiscal_year, agency, report_date, enacted_k in session.execute(
        statement
    ):
        if enacted_k is None:
            continue
        pb_cycle = int(fiscal_year) + 1
        reference_k = r1_aggregates.get((
            pb_cycle,
            int(fiscal_year),
            "discretionary",
            "CY Request",
            agency,
        ))
        ingested_k = float(enacted_k)
        residual_k, residual_pct, status = _compare(
            ingested_k, reference_k, tolerance_pct
        )
        if status == "no_reference":
            explanation = (
                f"PB{pb_cycle} R-1 CY Request is unavailable for {agency}"
            )
        elif status == "outside_tolerance":
            explanation = (
                f"latest DD 1416 enacted total ({report_date}) differs from "
                f"the PB{pb_cycle} R-1 CY Request; timing or reporting "
                "differences may explain the residual"
            )
        else:
            explanation = ""
        rows.append(TieOutRow(
            pb_cycle=pb_cycle,
            fiscal_year=int(fiscal_year),
            scope=AGENCY_TO_ACCOUNT.get(agency, agency),
            agency=agency,
            stream="discretionary",
            basis="CY Request",
            ingested_k=ingested_k,
            reference_k=reference_k,
            residual_k=residual_k,
            residual_pct=residual_pct,
            status=status,
            explanation=explanation,
        ))
    return rows


__all__ = [
    "Basis",
    "Stream",
    "TieOutRow",
    "TieStatus",
    "tie_out_dd1416",
    "tie_out_r1",
]
