"""Parse DD 1416 RDT&E quarterly execution workbooks by header meaning."""

from __future__ import annotations

import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl

from acquisition.dd1416_downloader import metadata_from_url


class DD1416ParseError(ValueError):
    """Raised when a workbook changes shape or fails its accounting identity."""


_SPACE_RE = re.compile(r"\s+")
_BA_RE = re.compile(r"\bBA\s*0*(\d+)\s*:", re.IGNORECASE)
_FY_PAIR_RE = re.compile(r"\b(\d{4})\s*[-–]\s*(\d{4})\b")

FIELD_MATCHERS = {
    "line_number": lambda value: value in {"bli#", "bli #"},
    "pe_number": lambda value: value == "bli",
    "program_title": lambda value: value.startswith("bli title"),
    "request": lambda value: "president's budget request" in value,
    "enacted": lambda value: "enacted appropriation" in value,
    "statutory": lambda value: "adjustments required by statute" in value,
    "suppl_resc_seq": lambda value: (
        "suppls" in value and "rescissions" in value
    ),
    "other": lambda value: (
        (value.startswith("other:") and "cancel" in value)
        or "cancelled account adjustments" in value
    ),
    "above": lambda value: "above threshold reprog" in value,
    "below": lambda value: "below threshold reprog" in value,
    "net": lambda value: value == "net",
}


def _header(value: object) -> str:
    if value is None:
        return ""
    return _SPACE_RE.sub(" ", str(value).replace("\n", " ")).strip().lower()


def _decimal(value: object) -> Decimal:
    """Normalize the mixed numeric formats used by the official sheets."""
    if value is None:
        return Decimal(0)
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    raw = str(value).strip()
    if not raw or raw in {"-", "—", "–"}:
        return Decimal(0)
    negative = raw.startswith("(") and raw.endswith(")")
    raw = raw.strip("()").replace(",", "").replace("$", "")
    try:
        number = Decimal(raw)
    except InvalidOperation as exc:
        raise DD1416ParseError(f"Invalid dollar value {value!r}") from exc
    return -number if negative else number


def _column_map(values: tuple[object, ...]) -> dict[str, int] | None:
    normalized = [_header(value) for value in values]
    if "bli" not in normalized or not any(
        "president's budget request" in value for value in normalized
    ):
        return None
    mapping: dict[str, int] = {}
    for field, matcher in FIELD_MATCHERS.items():
        indexes = [i for i, value in enumerate(normalized) if matcher(value)]
        if field == "line_number" and not indexes:
            # Early reports have PE/BLI identifiers but no separate line
            # number. This field is useful metadata, not part of the amount
            # reconciliation or PE join.
            continue
        if not indexes:
            raise DD1416ParseError(
                f"Expected at least one {field!r} column, found 0"
            )
        # Some releases export merged header/value cells as repeated adjacent
        # columns. The leftmost copy is the canonical value.
        mapping[field] = indexes[0]
    return mapping


def parse_workbook(path: Path, metadata: dict | None = None) -> list[dict]:
    """Parse and reconcile every PE row from one DD 1416 workbook."""
    path = Path(path)
    metadata = metadata or metadata_from_url(path.name)
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if len(workbook.sheetnames) != 1:
        raise DD1416ParseError(
            f"Expected one worksheet in {path.name}, found {workbook.sheetnames}"
        )

    sheet = workbook[workbook.sheetnames[0]]
    columns: dict[str, int] | None = None
    budget_activity: int | None = None
    parsed: list[dict] = []
    header_count = 0
    # Current reports state "In Dollars"; early reports state "Dollars in
    # Thousands". The database contract is always thousands of dollars.
    unit_divisor = Decimal(1000)

    try:
        for row_number, cells in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = tuple(cells)
            row_text = " ".join(_header(value) for value in values if value is not None)
            if "dollars in thousands" in row_text:
                unit_divisor = Decimal(1)
            elif re.search(r"\bin dollars\b", row_text):
                unit_divisor = Decimal(1000)
            ba_match = _BA_RE.search(row_text)
            if ba_match:
                budget_activity = int(ba_match.group(1))

            mapped = _column_map(values)
            if mapped is not None:
                columns = mapped
                header_count += 1
                continue
            if columns is None:
                continue

            first = _header(values[0] if values else None)
            if first.startswith("total ba"):
                continue
            fy_match = _FY_PAIR_RE.search(str(values[0] or ""))
            if not fy_match:
                continue

            fy_start, fy_end = map(int, fy_match.groups())
            pe_number = str(values[columns["pe_number"]] or "").strip()
            if not pe_number or pe_number.lower() == "bli":
                continue
            program_title = str(values[columns["program_title"]] or "").strip()
            if not program_title or program_title.upper().startswith("TOTAL BA"):
                continue

            dollars = {
                field: _decimal(values[columns[field]])
                for field in (
                    "request", "enacted", "statutory", "suppl_resc_seq",
                    "other", "above", "below", "net",
                )
            }
            calculated_net = sum(
                dollars[field]
                for field in (
                    "enacted", "statutory", "suppl_resc_seq", "other",
                    "above", "below",
                )
            )
            residual = calculated_net - dollars["net"]
            if abs(residual) > Decimal("1"):
                raise DD1416ParseError(
                    f"{path.name} row {row_number} PE {pe_number}: "
                    f"execution columns miss Net by ${residual}"
                )

            line_value = (
                values[columns["line_number"]]
                if "line_number" in columns else None
            )
            parsed.append({
                "pe_number": pe_number,
                "agency": metadata["agency"],
                "appropriation": metadata.get("appropriation", "RDTE"),
                "fy_start": fy_start,
                "fy_end": fy_end,
                "report_date": metadata["report_date"],
                "line_number": (
                    str(line_value).strip() if line_value is not None else None
                ),
                "program_title": program_title,
                "budget_activity": budget_activity,
                "request_k": float(dollars["request"] / unit_divisor),
                "enacted_k": float(dollars["enacted"] / unit_divisor),
                "statutory_adj_k": float(dollars["statutory"] / unit_divisor),
                "suppl_resc_seq_k": float(dollars["suppl_resc_seq"] / unit_divisor),
                "other_adj_k": float(dollars["other"] / unit_divisor),
                "above_threshold_reprog_k": float(dollars["above"] / unit_divisor),
                "below_threshold_reprog_k": float(dollars["below"] / unit_divisor),
                "net_k": float(dollars["net"] / unit_divisor),
                "source_file": path.name,
                "source_row": row_number,
            })
    finally:
        workbook.close()

    if header_count == 0:
        raise DD1416ParseError(f"No DD 1416 header found in {path.name}")
    if not parsed:
        raise DD1416ParseError(f"No PE rows found in {path.name}")
    if any(abs(row["enacted_k"]) > 1e8 for row in parsed):
        raise DD1416ParseError(
            f"Implausible enacted amount in {path.name}; dollars-to-thousands failed"
        )
    return parsed
