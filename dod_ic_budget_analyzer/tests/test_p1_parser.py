"""Tests for the official Comptroller P-1 workbook parser."""

from pathlib import Path
import re
import sys
import tempfile
import unittest

import openpyxl

APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

from parsing.xlsx_ingest import P1_FUNDING_TYPES, P1Record, parse_p1


FIXTURE = Path(__file__).parent / "fixtures" / "p1_fy2027_sample.xlsx"

EXPECTED_RECORDS = [
    P1Record(
        "9670A00005", "MQ-1 UAV", "Army", "Aircraft Procurement, Army",
        1, "1", "10", "Fixed Wing", "A", "Weapon System Cost",
        2025, "PY Actual", 240000.0, 8.0, "Add",
    ),
    P1Record(
        "9670A00005", "MQ-1 UAV", "Army", "Aircraft Procurement, Army",
        1, "1", "10", "Fixed Wing", "A", "Weapon System Cost",
        2026, "CY Request", 240000.0, 8.0, "Add",
    ),
    P1Record(
        "9672A00510", "Future UAS Family", "Army",
        "Aircraft Procurement, Army", 1, "2", "10", "Fixed Wing", "A",
        "Weapon System Cost", 2025, "PY Actual", 57902.0, None, "Add",
    ),
    P1Record(
        "9672A00510", "Future UAS Family", "Army",
        "Aircraft Procurement, Army", 1, "2", "10", "Fixed Wing", "A",
        "Weapon System Cost", 2026, "CY Request", 71459.0, None, "Add",
    ),
    P1Record(
        "9672A00510", "Future UAS Family", "Army",
        "Aircraft Procurement, Army", 1, "2", "10", "Fixed Wing", "A",
        "Weapon System Cost", 2026, "CY Mandatory", 26000.0, None, "Add",
    ),
    P1Record(
        "6761A12000", "Future Vertical Lift Family of Systems", "Army",
        "Aircraft Procurement, Army", 1, "9", "20", "Rotary", "C",
        "C (FY 2027 for FY 2028) (M)", 2027, "BY Request", 127217.0,
        None, "Non-Add",
    ),
    P1Record(
        "8046C20200", "TERMINAL HIGH ALTITUDE AREA DEFENSE (THAAD)",
        "Army", "Missile Procurement, Army", 2, "2", "10",
        "Surface-to-Air Missile System", "A", "Weapon System Cost", 2027,
        "BY Request", 907162.0, 27.0, "Add",
    ),
    P1Record(
        "8046C20200", "TERMINAL HIGH ALTITUDE AREA DEFENSE (THAAD)",
        "Army", "Missile Procurement, Army", 2, "2", "10",
        "Surface-to-Air Missile System", "A", "Weapon System Cost", 2027,
        "BY Mandatory", 10528043.0, 830.0, "Add",
    ),
    P1Record(
        "0122TA0600", "Information System Security Program-ISSP", "Army",
        "Other Procurement, Army", 2, "33", "64", "Information Security",
        "A", "Weapon System Cost", 2025, "PY Actual", 34037.0, None, "Add",
    ),
    P1Record(
        "0122TA0600", "Information System Security Program-ISSP", "Army",
        "Other Procurement, Army", 2, "33", "64", "Information Security",
        "A", "Weapon System Cost", 2025, "PY Mandatory", 1500.0, None,
        "Add",
    ),
    P1Record(
        "0122TA0600", "Information System Security Program-ISSP", "Army",
        "Other Procurement, Army", 2, "33", "64", "Information Security",
        "A", "Weapon System Cost", 2026, "CY Request", 826.0, None, "Add",
    ),
    P1Record(
        "0122TA0600", "Information System Security Program-ISSP", "Army",
        "Other Procurement, Army", 2, "33", "64", "Information Security",
        "A", "Weapon System Cost", 2027, "BY Request", 853.0, None, "Add",
    ),
    P1Record(
        "3050", "Ship Communications Automation", "Navy",
        "Other Procurement, Navy", 2, "74", "11", "Shipboard Communications",
        "A", "Weapon System Cost", 2025, "PY Actual", 103546.0, 0.0, "Add",
    ),
    P1Record(
        "3050", "Ship Communications Automation", "Navy",
        "Other Procurement, Navy", 2, "74", "11", "Shipboard Communications",
        "A", "Weapon System Cost", 2025, "PY Mandatory", 23775.0, 0.0,
        "Add",
    ),
    P1Record(
        "3050", "Ship Communications Automation", "Navy",
        "Other Procurement, Navy", 2, "74", "11", "Shipboard Communications",
        "A", "Weapon System Cost", 2026, "CY Request", 162075.0, 0.0,
        "Add",
    ),
    P1Record(
        "3050", "Ship Communications Automation", "Navy",
        "Other Procurement, Navy", 2, "74", "11", "Shipboard Communications",
        "A", "Weapon System Cost", 2027, "BY Request", 156605.0, 0.0,
        "Add",
    ),
    P1Record(
        "9670A00005", "MQ-1 UAV", "Army", "Aircraft Procurement, Army",
        1, "1", "10", "Fixed Wing", "B",
        "Less: Advance Procurement (PY)", 2026, "CY Request", -40000.0,
        None, "Add",
    ),
    P1Record(
        "9670A00005", "MQ-1 UAV", "Army", "Aircraft Procurement, Army",
        5, "44", "02", "Tactical Aircraft", "A", "Weapon System Cost",
        2026, "CY Request", 50000.0, 2.0, "Add",
    ),
]


class P1ParserTests(unittest.TestCase):
    def test_fixture_parses_to_exact_records(self):
        records = list(parse_p1(FIXTURE, pb_cycle=2027))

        self.assertEqual(records, EXPECTED_RECORDS)
        self.assertEqual(
            {record.funding_type for record in records},
            set(P1_FUNDING_TYPES),
        )
        self.assertTrue(
            any(record.add_non_add == "Non-Add" for record in records)
        )
        future_uas = [
            record for record in records if record.bli == "9672A00510"
        ]
        self.assertEqual(
            [(record.funding_type, record.amount_thousands)
             for record in future_uas],
            [
                ("PY Actual", 57902.0),
                ("CY Request", 71459.0),
                ("CY Mandatory", 26000.0),
            ],
        )

    def test_fixture_has_distinct_full_row_identities(self):
        records = [
            record for record in parse_p1(FIXTURE, pb_cycle=2027)
            if record.add_non_add == "Add"
        ]
        keys = [
            (
                record.bli,
                record.agency,
                record.appropriation,
                record.budget_activity,
                record.line_number,
                record.cost_type,
                record.cost_type_title,
                record.fiscal_year,
                record.funding_type,
                2027,
                FIXTURE.name,
            )
            for record in records
        ]
        self.assertEqual(len(keys), len(set(keys)))

        mq1_records = [
            record for record in records if record.bli == "9670A00005"
        ]
        self.assertEqual(
            {record.budget_activity for record in mq1_records}, {1, 5}
        )
        self.assertEqual(
            {record.cost_type for record in mq1_records}, {"A", "B"}
        )

    def test_base_columns_are_mapped_by_header_text(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            reordered = Path(temp_dir) / "reordered.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            for left_column, right_column in ((4, 9), (6, 11), (7, 12)):
                for row_number in range(1, worksheet.max_row + 1):
                    left = worksheet.cell(row_number, left_column).value
                    right = worksheet.cell(row_number, right_column).value
                    worksheet.cell(row_number, left_column).value = right
                    worksheet.cell(row_number, right_column).value = left
            workbook.save(reordered)
            workbook.close()

            self.assertEqual(
                list(parse_p1(reordered, pb_cycle=2027)), EXPECTED_RECORDS
            )

    def test_optional_mandatory_streams_and_disc_request_header(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            discretionary_only = Path(temp_dir) / "discretionary_only.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            for column in (16, 17, 22, 23, 28, 29):
                worksheet.cell(2, column).value = None
            for column in (26, 27):
                worksheet.cell(2, column).value = (
                    worksheet.cell(2, column).value.replace(
                        "Discretionary Request", "Disc Request"
                    )
                )
            workbook.save(discretionary_only)
            workbook.close()

            expected = [
                record for record in EXPECTED_RECORDS
                if not record.funding_type.endswith(" Mandatory")
            ]
            self.assertEqual(
                list(parse_p1(discretionary_only, pb_cycle=2027)), expected
            )

    def test_incomplete_scenario_pair_names_missing_header(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            malformed = Path(temp_dir) / "missing_quantity.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            worksheet.cell(2, 28).value = None
            workbook.save(malformed)
            workbook.close()

            with self.assertRaisesRegex(
                ValueError, "FY 2027 Mandatory Quantity"
            ):
                list(parse_p1(malformed, pb_cycle=2027))

    def test_missing_identity_header_is_named(self):
        headers = (
            (6, "Line Number"),
            (7, "BSA"),
            (8, "Budget SubActivity (BSA) Title"),
            (11, "Cost Type"),
            (12, "Cost Type Title"),
        )
        for column, header in headers:
            with (
                self.subTest(header=header),
                tempfile.TemporaryDirectory() as temp_dir,
            ):
                malformed = Path(temp_dir) / "missing_identity_header.xlsx"
                workbook = openpyxl.load_workbook(FIXTURE)
                worksheet = workbook["Exhibit P-1"]
                worksheet.cell(2, column).value = None
                workbook.save(malformed)
                workbook.close()

                with self.assertRaisesRegex(ValueError, re.escape(header)):
                    list(parse_p1(malformed, pb_cycle=2027))

    def test_missing_budget_activity_names_row(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            malformed = Path(temp_dir) / "missing_budget_activity.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            worksheet.cell(3, 4).value = None
            workbook.save(malformed)
            workbook.close()

            with self.assertRaisesRegex(
                ValueError, "Missing Budget Activity at row 3"
            ):
                list(parse_p1(malformed, pb_cycle=2027))

    def test_blank_identity_text_is_empty_string(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            blank_identity = Path(temp_dir) / "blank_identity.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            for column in (6, 7, 8, 11, 12):
                worksheet.cell(3, column).value = None
            workbook.save(blank_identity)
            workbook.close()

            record = next(parse_p1(blank_identity, pb_cycle=2027))
            self.assertEqual(
                (
                    record.line_number,
                    record.bsa,
                    record.bsa_title,
                    record.cost_type,
                    record.cost_type_title,
                ),
                ("", "", "", "", ""),
            )

    def test_missing_header_is_named(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            malformed = Path(temp_dir) / "malformed.xlsx"
            workbook = openpyxl.load_workbook(FIXTURE)
            worksheet = workbook["Exhibit P-1"]
            worksheet.cell(2, 10).value = None
            workbook.save(malformed)
            workbook.close()

            with self.assertRaisesRegex(
                ValueError, r"Budget Line Item \(BLI\) Title"
            ):
                list(parse_p1(malformed, pb_cycle=2027))


if __name__ == "__main__":
    unittest.main()
