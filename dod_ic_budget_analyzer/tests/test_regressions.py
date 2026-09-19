import sys
import tempfile
import unittest
import gzip
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

import pandas as pd
import polars as pl
import openpyxl
from sqlalchemy import create_engine, func, inspect, select, text
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

APP_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(APP_DIR))

import config
from analysis.ai_budget import LedgerUnavailable, SpendLedger, budget_guard
from analysis.gap_analyzer import GapAnalyzer
from analysis.deflators import apply_deflator, convert_amount
from analysis.execution_view import ExecutionView
from analysis.program_linker import ProgramLinker
from analysis.provenance import (
    csv_with_provenance,
    funding_sources,
    procurement_sources,
)
from analysis.rhetoric_tracker import align_rhetoric_funding
from analysis.spending_explorer import SpendingExplorer
from analysis.trend_tracker import TrendTracker
from analysis.user_identity import streamlit_user_id
from acquisition.dd1416_downloader import metadata_from_url
from parsing.dd1416_parser import DD1416ParseError, parse_workbook
from storage.db import (
    Base, FACT_TYPES, FundingLine, NarrativeFact, PEExecution, PENarrative,
    ProcurementLine, ProgramElement, SourceDocument,
    _ensure_schema_compatibility, get_engine, narrative_fact_hash,
)
from storage.ingest_r1 import R1Ingestor
from storage.build_archive import build_archive


class ArchiveRegressionTests(unittest.TestCase):
    def test_archive_round_trip_is_atomic_and_exact(self):
        with tempfile.TemporaryDirectory() as tmp:
            processed = Path(tmp)
            database = processed / "test.db"
            archive = processed / "test.db.gz"
            payload = (b"SQLite test payload\x00" * 10_000)
            database.write_bytes(payload)
            with patch("storage.build_archive.config.PROCESSED_DIR", processed):
                raw_size, archive_size = build_archive(database, archive)

            self.assertEqual(raw_size, len(payload))
            self.assertLess(archive_size, raw_size)
            with gzip.open(archive, "rb") as packaged:
                self.assertEqual(packaged.read(), payload)
            self.assertFalse((processed / "test.db.gz.tmp").exists())


class FundingRegressionTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        doc = SourceDocument(
            filename="seed.xlsx", document_type="R1", publication_year=2026
        )
        pe = ProgramElement(
            source_document=doc,
            pe_number="0207423F",
            program_name="Advanced Communications Systems",
            agency="Air Force",
        )
        self.session.add(pe)
        self.session.flush()
        self.pe_id = pe.id
        self.session.add_all([
            FundingLine(
                program_element_id=pe.id,
                fiscal_year=2024,
                funding_type="PY Actual",
                amount_thousands=11_264,
            ),
            FundingLine(
                program_element_id=pe.id,
                fiscal_year=2024,
                funding_type="PY Actual",
                amount_thousands=12_693,
            ),
            FundingLine(
                program_element_id=pe.id,
                fiscal_year=2024,
                funding_type="PY Mandatory",
                amount_thousands=99_999,
            ),
        ])
        self.session.commit()

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    def test_program_history_sums_lines_and_excludes_mandatory_stream(self):
        history = TrendTracker(self.session).get_pe_history(
            "0207423F", "Air Force"
        )
        self.assertEqual(history.height, 1)
        self.assertEqual(history.row(0, named=True)["amount_thousands"], 23_957)
        self.assertEqual(history.row(0, named=True)["basis"], "Actual")

    def test_gap_analyzer_sums_same_basis_before_selecting_it(self):
        funding = GapAnalyzer(self.session)._fetch_funding_data([self.pe_id])
        self.assertEqual(funding.height, 1)
        self.assertEqual(funding.row(0, named=True)["amount_thousands"], 23_957)

    def test_rdte_deflator_spot_check_and_conversion(self):
        self.assertAlmostEqual(convert_amount(84.30, 2020), 100.0, places=6)
        frame = pd.DataFrame({"fiscal_year": [2020, 2025], "amount": [84.3, 100.0]})
        converted = apply_deflator(frame, amount_column="amount")
        self.assertEqual([round(value, 6) for value in converted], [100.0, 100.0])


class ProcurementSchemaRegressionTests(unittest.TestCase):
    def test_compatibility_migrates_and_enforces_row_identity(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        ProcurementLine.__table__.drop(engine)
        with engine.begin() as connection:
            connection.execute(text("""
                CREATE TABLE procurement_lines (
                    id INTEGER NOT NULL PRIMARY KEY AUTOINCREMENT,
                    source_document_id INTEGER NOT NULL
                        REFERENCES source_documents(id),
                    bli VARCHAR(50) NOT NULL,
                    line_item_title VARCHAR(500) NOT NULL,
                    agency VARCHAR(100) NOT NULL,
                    appropriation VARCHAR(100) NOT NULL,
                    budget_activity INTEGER,
                    fiscal_year INTEGER NOT NULL,
                    funding_type VARCHAR(50) NOT NULL,
                    amount_thousands FLOAT NOT NULL,
                    quantity FLOAT,
                    pb_cycle INTEGER NOT NULL,
                    content_hash VARCHAR(64) NOT NULL,
                    ingested_at DATETIME NOT NULL,
                    UNIQUE (
                        bli, agency, appropriation, fiscal_year,
                        funding_type, pb_cycle, source_document_id
                    )
                )
            """))

        _ensure_schema_compatibility(engine)
        schema = inspect(engine)
        columns = {
            column["name"]: column
            for column in schema.get_columns("procurement_lines")
        }
        for column_name in (
            "line_number",
            "bsa",
            "bsa_title",
            "cost_type",
            "cost_type_title",
        ):
            self.assertIn(column_name, columns)
            self.assertFalse(columns[column_name]["nullable"])

        expected_unique = (
            "bli",
            "agency",
            "appropriation",
            "budget_activity",
            "line_number",
            "cost_type",
            "cost_type_title",
            "fiscal_year",
            "funding_type",
            "pb_cycle",
            "source_document_id",
        )
        legacy_unique = (
            "bli",
            "agency",
            "appropriation",
            "fiscal_year",
            "funding_type",
            "pb_cycle",
            "source_document_id",
        )
        unique_constraints = {
            tuple(constraint["column_names"])
            for constraint in schema.get_unique_constraints(
                "procurement_lines"
            )
        }
        self.assertIn(expected_unique, unique_constraints)
        self.assertNotIn(legacy_unique, unique_constraints)

        with Session(engine) as session:
            source = SourceDocument(
                filename="fy2027_p1.xlsx",
                document_type="P1",
                publication_year=2027,
            )
            session.add(source)
            session.flush()

            common = {
                "source_document_id": source.id,
                "bli": "F015EX",
                "line_item_title": "F-15EX",
                "agency": "Air Force",
                "appropriation": "Aircraft Procurement, Air Force",
                "fiscal_year": 2027,
                "funding_type": "BY Request",
                "amount_thousands": 1_250_000.0,
                "quantity": None,
                "pb_cycle": 2027,
            }
            first = ProcurementLine(
                **common,
                budget_activity=1,
                line_number="7",
                bsa="03",
                bsa_title="Tactical Forces",
                cost_type="A",
                cost_type_title="Weapon System Cost",
                content_hash="a" * 64,
            )
            second = ProcurementLine(
                **common,
                budget_activity=1,
                line_number="7",
                bsa="03",
                bsa_title="Tactical Forces",
                cost_type="B",
                cost_type_title="Less: Advance Procurement (PY)",
                content_hash="b" * 64,
            )
            third = ProcurementLine(
                **common,
                budget_activity=5,
                line_number="44",
                bsa="02",
                bsa_title="Tactical Aircraft",
                cost_type="A",
                cost_type_title="Weapon System Cost",
                content_hash="c" * 64,
            )
            session.add_all((first, second, third))
            session.commit()

            stored = session.scalars(
                select(ProcurementLine).order_by(ProcurementLine.id)
            ).all()
            self.assertEqual(len(stored), 3)
            self.assertEqual(
                {
                    (
                        row.budget_activity,
                        row.line_number,
                        row.cost_type,
                        row.cost_type_title,
                    )
                    for row in stored
                },
                {
                    (1, "7", "A", "Weapon System Cost"),
                    (1, "7", "B", "Less: Advance Procurement (PY)"),
                    (5, "44", "A", "Weapon System Cost"),
                },
            )

            session.add(ProcurementLine(
                **common,
                budget_activity=1,
                line_number="7",
                bsa="different-but-not-keyed",
                bsa_title="Different but not keyed",
                cost_type="A",
                cost_type_title="Weapon System Cost",
                content_hash="d" * 64,
            ))
            with self.assertRaises(IntegrityError):
                session.commit()
            session.rollback()
        engine.dispose()


class NarrativeFactSchemaRegressionTests(unittest.TestCase):
    def setUp(self):
        self.engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(self.engine)
        self.session = Session(self.engine)
        self.description = (
            "The program hired Lockheed Martin. Testing follows next year."
        )
        self.sentence = "The program hired Lockheed Martin."
        self.narrative = PENarrative(
            pe_number="0600001A",
            agency="Army",
            fiscal_year=2027,
            project_number="001",
            project_title="Example project",
            description=self.description,
            source_file="army_fy2027_example.pdf",
        )
        self.session.add(self.narrative)
        self.session.flush()

    def tearDown(self):
        self.session.close()
        self.engine.dispose()

    def _fact(self) -> NarrativeFact:
        char_start = self.description.index(self.sentence)
        char_end = char_start + len(self.sentence)
        return NarrativeFact(
            narrative_table="pe_narratives",
            narrative_id=self.narrative.id,
            pe_number=self.narrative.pe_number,
            agency=self.narrative.agency,
            fiscal_year=self.narrative.fiscal_year,
            fact_type="contractor",
            value="Lockheed Martin",
            sentence=self.sentence,
            char_start=char_start,
            char_end=char_end,
            model=config.GEMINI_MODEL,
            content_hash=narrative_fact_hash(
                "pe_narratives",
                self.narrative.id,
                "contractor",
                "Lockheed Martin",
                char_start,
                char_end,
            ),
        )

    def test_compatibility_creates_table_on_existing_database(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "existing.db"
            engine = create_engine(f"sqlite:///{path.as_posix()}")
            Base.metadata.create_all(
                engine,
                tables=[
                    table for table in Base.metadata.sorted_tables
                    if table.name != "narrative_facts"
                ],
            )
            self.assertFalse(inspect(engine).has_table("narrative_facts"))
            engine.dispose()

            compatible = get_engine(f"sqlite:///{path.as_posix()}")
            try:
                self.assertTrue(
                    inspect(compatible).has_table("narrative_facts")
                )
            finally:
                compatible.dispose()

    def test_fact_round_trip_preserves_source_offsets(self):
        fact = self._fact()
        self.session.add(fact)
        self.session.commit()
        self.session.expire_all()

        stored = self.session.scalar(select(NarrativeFact))
        self.assertIsNotNone(stored)
        self.assertIsInstance(stored.id, int)
        self.assertEqual(stored.narrative_table, "pe_narratives")
        self.assertEqual(stored.narrative_id, self.narrative.id)
        self.assertEqual(stored.pe_number, "0600001A")
        self.assertEqual(stored.agency, "Army")
        self.assertEqual(stored.fiscal_year, 2027)
        self.assertEqual(stored.fact_type, "contractor")
        self.assertEqual(stored.value, "Lockheed Martin")
        self.assertEqual(stored.sentence, self.sentence)
        self.assertEqual(stored.char_start, 0)
        self.assertEqual(stored.char_end, len(self.sentence))
        self.assertEqual(stored.model, config.GEMINI_MODEL)
        self.assertEqual(stored.content_hash, fact.content_hash)
        self.assertIsNotNone(stored.extracted_at)
        self.assertEqual(
            self.description[stored.char_start:stored.char_end],
            stored.sentence,
        )

    def test_duplicate_content_hash_is_rejected(self):
        first = self._fact()
        self.session.add(first)
        self.session.commit()

        duplicate = self._fact()
        self.session.add(duplicate)
        with self.assertRaises(IntegrityError):
            self.session.commit()
        self.session.rollback()

    def test_fact_types_are_exact(self):
        self.assertEqual(
            set(FACT_TYPES),
            {"contractor", "transition", "test_event", "location"},
        )


class IngestionRegressionTests(unittest.TestCase):
    def test_ingest_is_idempotent_and_retains_mandatory_streams(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        session = Session(engine)
        try:
            with tempfile.TemporaryDirectory() as tmp:
                parquet = Path(tmp) / "r1.parquet"
                pl.DataFrame([{
                    "source_file": "fy2027_r1.xlsx",
                    "fiscal_year": 2027,
                    "pe_number": "0600001A",
                    "component": "Army",
                    "act_code": "04",
                    "line_no": 1,
                    "pe_title": "Test Program",
                    "py_amount": 10.0,
                    "cy_amount": 20.0,
                    "by_amount": 30.0,
                    "py_mandatory_amount": 1.0,
                    "cy_mandatory_amount": 2.0,
                    "by_mandatory_amount": 3.0,
                }]).write_parquet(parquet)

                ingestor = R1Ingestor(session)
                ingestor.ingest_parquet(str(parquet))
                ingestor.ingest_parquet(str(parquet))

            count = session.execute(select(func.count(FundingLine.id))).scalar_one()
            types = set(session.execute(select(FundingLine.funding_type)).scalars())
            self.assertEqual(count, 6)
            self.assertEqual(
                types,
                {
                    "PY Actual", "CY Request", "BY Request",
                    "PY Mandatory", "CY Mandatory", "BY Mandatory",
                },
            )
            rows = session.execute(select(FundingLine)).scalars().all()
            self.assertTrue(all(row.pb_cycle == 2027 for row in rows))
            self.assertTrue(all(row.source_document_id is not None for row in rows))
        finally:
            session.close()
            engine.dispose()


class ProvenanceRegressionTests(unittest.TestCase):
    def test_funding_export_names_its_source(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        with Session(engine) as session:
            doc = SourceDocument(
                filename="fy2027_r1.xlsx",
                document_type="R1",
                publication_year=2027,
                source_url="https://example.test/fy2027_r1.xlsx",
            )
            pe = ProgramElement(
                source_document=doc,
                pe_number="0600001A",
                program_name="Test Program",
                agency="Army",
            )
            session.add(pe)
            session.flush()
            session.add(FundingLine(
                program_element=pe,
                source_document=doc,
                pb_cycle=2027,
                fiscal_year=2027,
                funding_type="BY Request",
                amount_thousands=123.0,
            ))
            session.commit()

            sources = funding_sources(
                session, pe_numbers=["0600001A"], agencies=["Army"]
            )
            exported = csv_with_provenance(
                pd.DataFrame([{"FY": 2027, "$K": 123.0}]), sources
            ).decode("utf-8-sig")

        self.assertEqual(len(sources), 1)
        self.assertIn("fy2027_r1.xlsx", exported)
        self.assertIn("https://example.test/fy2027_r1.xlsx", exported)

    def test_procurement_export_names_its_p1_source(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        with Session(engine) as session:
            doc = SourceDocument(
                filename="fy2027_p1.xlsx",
                document_type="P1",
                publication_year=2027,
                source_url="https://example.test/fy2027_p1.xlsx",
            )
            session.add(doc)
            session.flush()
            session.add(ProcurementLine(
                source_document_id=doc.id,
                bli="F015EX",
                line_item_title="F-15EX",
                agency="Air Force",
                appropriation="Aircraft Procurement, Air Force",
                budget_activity=1,
                line_number="1",
                bsa="01",
                bsa_title="Aircraft",
                cost_type="C",
                cost_type_title="Cost",
                fiscal_year=2027,
                funding_type="BY Request",
                amount_thousands=123.0,
                quantity=1.0,
                pb_cycle=2027,
                content_hash="synthetic-p1-provenance",
            ))
            session.commit()

            sources = procurement_sources(
                session,
                blis=["F015EX"],
                appropriations=["Aircraft Procurement, Air Force"],
            )
            exported = csv_with_provenance(
                pd.DataFrame([{"BLI": "F015EX"}]), sources
            ).decode("utf-8-sig")

        self.assertEqual(len(sources), 1)
        self.assertEqual(sources[0]["document_type"], "P1")
        self.assertIn("fy2027_p1.xlsx", exported)


class DD1416RegressionTests(unittest.TestCase):
    HEADERS = {
        "line_number": "BLI#",
        "pe_number": "BLI",
        "program_title": "BLI TITLE",
        "request": "President's Budget Request",
        "enacted": (
            "Enacted Appropriation (Includes Distribution of "
            "Congressional Adjustments/1)"
        ),
        "statutory": "Adjustments Required by Statute /2",
        "suppl": "Suppls/Collections/Rescissions/ Sequestration /3",
        "other": "Other: Cancelled, Claims, Judgments /4",
        "above": "Above Threshold Reprog",
        "below": "Below Threshold Reprog",
        "net": "Net",
    }

    def _workbook(self, path: Path, columns: dict[str, int], bad_net=False):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "ORG-BA-BLI"
        sheet.cell(3, 1, "BA 01: BASIC RESEARCH")
        for key, title in self.HEADERS.items():
            sheet.cell(4, columns[key] + 1, title)
        values = {
            "line_number": "1",
            "pe_number": "0601102A",
            "program_title": "Defense Research Sciences",
            "request": 310_191_000,
            "enacted": 297_680_000,
            "statutory": "-3,988,648",
            "suppl": 0,
            "other": -31_050,
            "above": 0,
            "below": -3_196_000,
            "net": 1 if bad_net else 290_464_302,
        }
        sheet.cell(5, 1, "2025-2026")
        for key, value in values.items():
            sheet.cell(5, columns[key] + 1, value)
        sheet.cell(6, 1, "TOTAL BA 01")
        workbook.save(path)
        workbook.close()

    def test_component_column_offsets_map_to_same_semantics(self):
        defense_columns = {
            "line_number": 3, "pe_number": 4, "program_title": 5,
            "request": 7, "enacted": 9, "statutory": 11,
            "suppl": 13, "other": 15, "above": 17, "below": 19,
            "net": 21,
        }
        army_columns = {
            "line_number": 3, "pe_number": 4, "program_title": 7,
            "request": 10, "enacted": 12, "statutory": 14,
            "suppl": 16, "other": 19, "above": 23, "below": 27,
            "net": 29,
        }
        with tempfile.TemporaryDirectory() as tmp:
            defense = Path(tmp) / (
                "Defense_Wide_RDTE_FY_2025_2026_DD_1416_Qtrly_Rpt_"
                "03_31_2026.xlsx"
            )
            army = Path(tmp) / (
                "Army_RDTE_FY_2025_2026_DD_1416_Qtrly_Rpt_03_31_2026.xlsx"
            )
            self._workbook(defense, defense_columns)
            self._workbook(army, army_columns)
            defense_row = parse_workbook(defense)[0]
            army_row = parse_workbook(army)[0]

        for field in (
            "pe_number", "request_k", "enacted_k", "statutory_adj_k",
            "below_threshold_reprog_k", "net_k",
        ):
            self.assertEqual(defense_row[field], army_row[field])
        self.assertEqual(defense_row["net_k"], 290_464.302)

    def test_legacy_filename_order_is_discovered(self):
        item = metadata_from_url(
            "Army_FY_2013_2014_DD_1416_RDTE_Qtrly_Rpt_9_30_2013.xlsx"
        )
        self.assertEqual(item["agency"], "Army")
        self.assertEqual((item["fy_start"], item["fy_end"]), (2013, 2014))
        self.assertEqual(item["report_date"], "2013-09-30")

    def test_legacy_thousands_units_and_missing_line_number(self):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.cell(2, 1, "(Dollars in Thousands)")
        sheet.cell(3, 1, "BA 01: BASIC RESEARCH")
        legacy_headers = [
            "Period of Availability", "BLI", "BLI TITLE",
            "President's Budget Request", "Enacted Appropriation",
            "Adjustments Required by Statute /6", "Suppls/Rescissions",
            "Cancelled Account Adjustments", "Above Threshold Reprog",
            "Below Threshold Reprog", "Net",
        ]
        for column, title in enumerate(legacy_headers, start=1):
            sheet.cell(4, column, title)
        for column, value in enumerate([
            "2012-2013", "0601102F", "Defense Research Sciences", 364328,
            364328, -7847, 0, 0, 0, -8648, 347833,
        ], start=1):
            sheet.cell(5, column, value)
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / (
                "Air_Force_FY_2012_2013_DD_1416_RDTE_Qtrly_Rpt_"
                "12_31_2012.xlsx"
            )
            workbook.save(path)
            workbook.close()
            row = parse_workbook(path)[0]
        self.assertIsNone(row["line_number"])
        self.assertEqual(row["enacted_k"], 364328.0)
        self.assertEqual(row["net_k"], 347833.0)

    def test_reconciliation_failure_is_rejected(self):
        columns = {
            "line_number": 3, "pe_number": 4, "program_title": 5,
            "request": 7, "enacted": 9, "statutory": 11,
            "suppl": 13, "other": 15, "above": 17, "below": 19,
            "net": 21,
        }
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / (
                "Army_RDTE_FY_2025_2026_DD_1416_Qtrly_Rpt_03_31_2026.xlsx"
            )
            self._workbook(path, columns, bad_net=True)
            with self.assertRaises(DD1416ParseError):
                parse_workbook(path)

    def test_execution_view_keeps_latest_quarter(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        with Session(engine) as session:
            for report_date, net in (("2025-12-31", 90.0), ("2026-03-31", 95.0)):
                session.add(PEExecution(
                    pe_number="0601102A", agency="Army", appropriation="RDTE",
                    fy_start=2025, fy_end=2026, report_date=report_date,
                    program_title="Defense Research Sciences",
                    request_k=100.0, enacted_k=100.0, statutory_adj_k=0.0,
                    suppl_resc_seq_k=0.0, other_adj_k=0.0,
                    above_threshold_reprog_k=0.0,
                    below_threshold_reprog_k=net - 100.0, net_k=net,
                    source_file=f"{report_date}.xlsx",
                    content_hash=f"hash-{report_date}",
                ))
            session.commit()
            result = ExecutionView(session).latest_program_series(
                ["0601102A"], ["Army"]
            )
        self.assertEqual(result.height, 1)
        self.assertEqual(result.row(0, named=True)["net_k"], 95.0)

    def test_execution_view_filters_exact_pe_component_pairs(self):
        engine = create_engine("sqlite:///:memory:")
        Base.metadata.create_all(engine)
        with Session(engine) as session:
            for pe_number, agency, net in (
                ("0600001A", "Army", 10.0),
                ("0600002N", "Navy", 20.0),
                ("0600001A", "Navy", 1_000.0),
            ):
                session.add(PEExecution(
                    pe_number=pe_number, agency=agency, appropriation="RDTE",
                    fy_start=2025, fy_end=2026, report_date="2026-03-31",
                    program_title="Test", request_k=net, enacted_k=net,
                    statutory_adj_k=0.0, suppl_resc_seq_k=0.0,
                    other_adj_k=0.0, above_threshold_reprog_k=0.0,
                    below_threshold_reprog_k=0.0, net_k=net,
                    source_file=f"{pe_number}-{agency}.xlsx",
                    content_hash=f"{pe_number}-{agency}",
                ))
            session.commit()
            result = ExecutionView(session).latest_program_series(
                ["0600001A", "0600002N"], ["Army", "Navy"]
            )
        self.assertEqual(result.row(0, named=True)["net_k"], 30.0)


class MatchingRegressionTests(unittest.TestCase):
    def test_exact_pe_across_agencies_requires_review(self):
        candidates = [
            {
                "pe_id": 1,
                "name": "Space Weather",
                "pe_number": "0604002SF",
                "agency": "Air Force",
                "score": 1.0,
                "strategy": "PE_NUMBER",
            },
            {
                "pe_id": 2,
                "name": "Space Weather",
                "pe_number": "0604002SF",
                "agency": "Space Force",
                "score": 1.0,
                "strategy": "PE_NUMBER",
            },
        ]

        class ExactMatcher:
            def lookup_pe_number(self, _query):
                return candidates

        result = ProgramLinker(ExactMatcher()).link_query("PE 0604002SF")
        self.assertTrue(result["needs_review"])
        self.assertEqual(result["match_strategy"], "AMBIGUOUS")


class ExternalQueryRegressionTests(unittest.TestCase):
    def test_subaward_queries_are_scoped_to_dod(self):
        class FakeHTTP:
            def __init__(self):
                self.payloads = []

            def post(self, _path, payload):
                self.payloads.append(payload)
                return {"results": []}

        explorer = SpendingExplorer.__new__(SpendingExplorer)
        explorer.http = FakeHTTP()
        explorer.program_subawards("Test Program", 2024, 2025)

        self.assertEqual(len(explorer.http.payloads), 2)
        for payload in explorer.http.payloads:
            self.assertEqual(
                payload["filters"]["agencies"], explorer.DOD_FUNDING
            )


class GovernanceRegressionTests(unittest.TestCase):
    def test_budget_guard_fails_closed_when_ledger_is_unavailable(self):
        with patch.object(
            SpendLedger,
            "month_to_date",
            side_effect=LedgerUnavailable("locked"),
        ):
            decision = budget_guard("adjudicate", user_id="alice")
        self.assertFalse(decision.allowed)
        self.assertEqual(decision.reason, "metering")

    def test_anonymous_streamlit_sessions_do_not_share_identity(self):
        first = SimpleNamespace(
            user=SimpleNamespace(is_logged_in=False), session_state={}
        )
        second = SimpleNamespace(
            user=SimpleNamespace(is_logged_in=False), session_state={}
        )
        first_id = streamlit_user_id(first)
        self.assertEqual(streamlit_user_id(first), first_id)
        self.assertNotEqual(streamlit_user_id(second), first_id)


class StatisticsRegressionTests(unittest.TestCase):
    def test_alignment_reports_actual_overlap_count(self):
        signal = pd.DataFrame({
            "fiscal_year": range(2020, 2026),
            "mention_intensity": [1, 2, 3, 4, 5, 6],
            "positive_pct": [50] * 6,
            "negative_pct": [10] * 6,
            "stated_priority": [False] * 6,
        })
        funding = pd.DataFrame({
            "fiscal_year": [2021, 2022, 2024, 2025],
            "amount_thousands": [1_000, 2_000, 4_000, 5_000],
        })
        result = align_rhetoric_funding(signal, funding)
        self.assertEqual(result["alignment"]["n_years"], 4)


if __name__ == "__main__":
    unittest.main()
